import io
import json
import os
from datetime import datetime
from pathlib import Path

import pandas as pd
import streamlit as st
from docx import Document
from pypdf import PdfReader
from openai import OpenAI

st.set_page_config(page_title="Revisor de informes MIN328", page_icon="📋", layout="wide")

RUBRICA = [
    {"seccion":"1. Índice","maximo":0.25,"niveles":[
        (0.25,"Todas las secciones están incluidas y la paginación es correcta."),
        (0.00,"No presenta índice o no permite ubicar las secciones.")]},
    {"seccion":"2. Abstract (en inglés)","maximo":0.50,"niveles":[
        (0.50,"Resumen completo, claro y correctamente redactado en inglés."),
        (0.25,"Está presente, pero es incompleto, poco claro o presenta errores que dificultan parcialmente la comprensión."),
        (0.00,"No presenta abstract, no está en inglés o no corresponde al informe.")]},
    {"seccion":"3. Objetivos","maximo":0.25,"niveles":[
        (0.25,"Los objetivos son claros y están relacionados con la práctica."),
        (0.00,"No presenta objetivos o no guardan relación con la práctica.")]},
    {"seccion":"4. Descripción de los equipos para medir parámetros","maximo":0.25,"niveles":[
        (0.25,"Describe de forma pertinente los equipos utilizados."),
        (0.00,"No presenta la descripción o no identifica los equipos utilizados.")]},
    {"seccion":"5. Procedimiento realizado","maximo":0.25,"niveles":[
        (0.25,"Procedimiento comprensible, con esquema y observaciones; las figuras incluidas están tituladas y numeradas."),
        (0.00,"No presenta el procedimiento o la información es insuficiente para comprender el ensayo.")]},
    {"seccion":"6. Presentación de resultados","maximo":2.00,"niveles":[
        (2.00,"Resultados completos, ordenados y verificables; incluye valores, condiciones de trabajo y cálculos necesarios."),
        (1.75,"Resultados sólidos y casi completos; falta un detalle menor."),
        (1.50,"Resultados adecuados, pero falta un elemento importante o hay omisiones puntuales."),
        (1.25,"Presenta la mayoría de los resultados, con varias omisiones o cálculos incompletos."),
        (1.00,"Presentación parcial; faltan datos, condiciones o cálculos relevantes."),
        (0.75,"Incluye pocos resultados útiles y existen omisiones importantes."),
        (0.50,"Resultados muy limitados, desordenados o difíciles de verificar."),
        (0.25,"Incluye alguna información relacionada con resultados, pero no permite evaluarlos adecuadamente."),
        (0.00,"No presenta resultados.")]},
    {"seccion":"7. Discusión de resultados","maximo":2.00,"niveles":[
        (2.00,"Interpreta los resultados y presenta comparaciones pertinentes, organizadas y sustentadas con los estándares solicitados."),
        (1.75,"Discusión sólida; presenta una omisión menor en la comparación o interpretación."),
        (1.50,"Interpreta los principales resultados, pero parte de las comparaciones o justificaciones está incompleta."),
        (1.25,"Discusión parcial; faltan varias comparaciones o explicaciones."),
        (1.00,"Interpretación limitada y comparaciones incompletas."),
        (0.75,"Menciona resultados o estándares, pero explica poco su significado."),
        (0.50,"Discusión superficial, con comparaciones escasas o poco sustentadas."),
        (0.25,"Incluye algún comentario relacionado con los resultados, sin discusión suficiente."),
        (0.00,"No presenta discusión de resultados.")]},
    {"seccion":"8. Conclusiones","maximo":2.00,"niveles":[
        (2.00,"Las conclusiones responden a todos los objetivos, incluyen resultados relevantes y consideran el caso práctico."),
        (1.75,"Conclusiones claras y vinculadas a los objetivos, con una omisión menor."),
        (1.50,"Responden a la mayoría de los objetivos, pero falta precisión en algún resultado o en el caso práctico."),
        (1.25,"Varias conclusiones son generales o no se vinculan claramente con los objetivos."),
        (1.00,"Conclusiones parciales; varias no están sustentadas en los resultados."),
        (0.75,"Conclusiones poco específicas y débilmente relacionadas con los objetivos."),
        (0.50,"Incluye afirmaciones generales con escasa relación con los resultados."),
        (0.25,"Presenta una afirmación final relacionada con la práctica, pero no desarrolla conclusiones suficientes."),
        (0.00,"No presenta conclusiones.")]},
    {"seccion":"9. Bibliografía","maximo":0.50,"niveles":[
        (0.50,"Fuentes pertinentes y referencias completas."),
        (0.25,"Presenta referencias, pero son incompletas, poco claras o tienen errores importantes."),
        (0.00,"No presenta bibliografía.")]},
    {"seccion":"10. Anexos: caso práctico","maximo":2.00,"niveles":[
        (2.00,"Caso práctico desarrollado integralmente, con resultados, discusión y conclusiones coherentes."),
        (1.75,"Desarrollo completo en general, con una omisión menor."),
        (1.50,"Incluye las tres partes, pero una es incompleta."),
        (1.25,"Desarrollo parcial, con resultados o discusión insuficientes."),
        (1.00,"Caso práctico incompleto; faltan elementos relevantes."),
        (0.75,"Incluye parte del caso, con análisis limitado."),
        (0.50,"Desarrollo muy breve y poco sustentado."),
        (0.25,"Incluye información del caso, pero no desarrolla adecuadamente resultados, discusión y conclusiones."),
        (0.00,"No presenta el caso práctico.")]},
]
TOTAL = sum(r["maximo"] for r in RUBRICA)

def extract_docx(data):
    doc = Document(io.BytesIO(data))
    chunks = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            chunks.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(chunks)

def extract_pdf(data):
    reader = PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)

def extract_report(upload):
    if upload.name.lower().endswith(".docx"):
        return extract_docx(upload.getvalue())
    return extract_pdf(upload.getvalue())

def evaluate(client, model, filename, content, session_info):
    rubric_text = "\n".join(
        f'{r["seccion"]} (máximo {r["maximo"]:.2f}):\n' +
        "\n".join(f'{score:.2f} puntos: {desc}' for score, desc in r["niveles"])
        for r in RUBRICA
    )
    system = """Eres un asistente de apoyo a la evaluación de informes de laboratorio universitarios.
Aplica exactamente los descriptores oficiales proporcionados; no los sustituyas por criterios genéricos.
Evalúa solo la evidencia disponible. No inventes contenido ni supongas que figuras, paginación o imágenes cumplen si no pueden inspeccionarse.
Elige exclusivamente uno de los puntajes listados para cada sección. Si está ausente o no cumple, asigna 0 cuando así lo establezca la rúbrica.
La carátula no se puntúa. Puntaje máximo total: 10.00.
La evaluación es una propuesta que la docente debe revisar.
Devuelve JSON válido con:
{"observaciones_generales":"...", "secciones":[{"seccion":"nombre exacto","puntaje":0.0,"evidencia":"...","mejoras":"..."}], "posibles_alertas":["..."]}
Incluye todas las secciones en el orden recibido."""
    user = f"""Datos de la actividad: {session_info or 'No especificados'}.
Archivo: {filename}

RÚBRICA OFICIAL:
{rubric_text}

TEXTO EXTRAÍDO DEL INFORME:
--- INICIO ---
{content[:100000]}
--- FIN ---

Evalúa sección por sección con el nivel exacto de la rúbrica.
En Discusión de resultados, revisa la comparación con los LMP de Perú y los estándares de calidad de agua de Brasil, Canadá y Chile, según corresponda a la sesión. Si no se puede saber qué estándares correspondían, señálalo.
En Procedimiento, ten en cuenta esquema, observaciones y figuras tituladas/numeradas; si el texto extraído no permite verificar elementos visuales, indícalo como limitación.
En Conclusiones, verifica su vínculo con objetivos, resultados y caso práctico.
En Anexos: caso práctico, verifica resultados, discusión y conclusiones.
Incluye evidencia textual breve y específica. Advierte limitaciones de extracción de tablas, imágenes o paginación."""
    response = client.chat.completions.create(
        model=model, temperature=0.1, response_format={"type":"json_object"},
        messages=[{"role":"system","content":system},{"role":"user","content":user}]
    )
    raw = json.loads(response.choices[0].message.content)
    given = {x.get("seccion"): x for x in raw.get("secciones", [])}
    rows = []
    for rubric in RUBRICA:
        item = given.get(rubric["seccion"], {})
        valid = [score for score, _ in rubric["niveles"]]
        try:
            score = float(item.get("puntaje", 0))
        except (TypeError, ValueError):
            score = 0.0
        score = min(valid, key=lambda n: abs(n-score))
        rows.append({
            "Sección": rubric["seccion"], "Puntaje": score, "Máximo": rubric["maximo"],
            "Evidencia identificada": str(item.get("evidencia", "No se indicó evidencia.")),
            "Mejora sugerida": str(item.get("mejoras", "Revisar manualmente."))
        })
    return {
        "archivo": filename,
        "estudiante": Path(filename).stem,
        "puntaje_total": round(sum(r["Puntaje"] for r in rows), 2),
        "observaciones_generales": str(raw.get("observaciones_generales", "")),
        "secciones": rows,
        "alertas": raw.get("posibles_alertas", [])
    }

def make_excel(reviews):
    summary_rows, detail_rows = [], []
    for r in reviews:
        summary = {"Archivo": r["archivo"], "Estudiante (editable)": r.get("estudiante", Path(r["archivo"]).stem)}
        for section in r["secciones"]:
            summary[section["Sección"]] = section["Puntaje"]
        summary["Total / 10"] = round(sum(section["Puntaje"] for section in r["secciones"]), 2)
        summary["Observaciones generales"] = r["observaciones_generales"]
        summary_rows.append(summary)
        for section in r["secciones"]:
            detail_rows.append({
                "Archivo": r["archivo"], "Estudiante": r.get("estudiante", ""),
                "Sección": section["Sección"], "Puntaje": section["Puntaje"],
                "Máximo": section["Máximo"], "Evidencia": section["Evidencia identificada"],
                "Mejora sugerida": section["Mejora sugerida"]
            })
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame(summary_rows).to_excel(writer, index=False, sheet_name="Consolidado")
        pd.DataFrame(detail_rows).to_excel(writer, index=False, sheet_name="Detalle por sección")
        for sheet in writer.book.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for col in sheet.columns:
                max_len = min(max(len(str(cell.value or "")) for cell in col) + 2, 55)
                sheet.column_dimensions[col[0].column_letter].width = max(12, max_len)
    output.seek(0)
    return output.getvalue()

st.title("Revisor de informes de laboratorio — MIN328")
st.caption("Minería y Medio Ambiente · Laboratorio 2026-2 · Rúbrica oficial de 10 puntos")
st.info("Carga los informes de una clase. La IA propone puntajes y observaciones según la rúbrica oficial. Revisa la evidencia y confirma las notas antes de usarlas.")
with st.expander("Ver rúbrica incorporada", expanded=False):
    for r in RUBRICA:
        st.markdown(f'**{r["seccion"]} — máximo {r["maximo"]:.2f} puntos**')
        for score, desc in r["niveles"]:
            st.write(f"- {score:.2f}: {desc}")
    st.markdown("**Total máximo: 10.00 puntos. La carátula no se califica.**")

with st.sidebar:
    st.header("Configuración")
    api_key = st.text_input("Clave API de OpenAI", type="password", help="También puede configurarse como secreto OPENAI_API_KEY en el despliegue.")
    model = st.text_input("Modelo de IA", value="gpt-4.1-mini")
    curso = st.text_input("Curso / sesión (opcional)", value="MIN328 Minería y Medio Ambiente — Laboratorio 2026-2")
    st.caption("La API puede generar costos independientes de la suscripción a ChatGPT.")

uploads = st.file_uploader("Sube todos los informes de la clase", type=["docx","pdf"], accept_multiple_files=True)
if st.button("Evaluar informes con IA", type="primary", disabled=not uploads):
    try:
        secret_key = st.secrets.get("OPENAI_API_KEY", "")
    except Exception:
        secret_key = ""
    key = api_key or secret_key or os.getenv("OPENAI_API_KEY")
    if not key:
        st.error("Introduce la clave API de OpenAI en la barra lateral.")
    else:
        client = OpenAI(api_key=key)
        completed = []
        progress = st.progress(0)
        for i, up in enumerate(uploads):
            with st.spinner(f"Analizando {up.name}..."):
                try:
                    txt = extract_report(up)
                    if not txt.strip():
                        raise ValueError("No se pudo extraer texto. Si es PDF escaneado, aplica OCR primero.")
                    completed.append(evaluate(client, model, up.name, txt, curso))
                except Exception as e:
                    st.error(f"No se pudo evaluar {up.name}: {e}")
            progress.progress((i+1)/len(uploads))
        if completed:
            st.session_state["min328_reviews"] = completed

reviews = st.session_state.get("min328_reviews", [])
if reviews:
    st.divider()
    st.header("Consolidado de la clase")
    st.write("Puedes revisar y editar el nombre del estudiante y los puntajes propuestos. Los puntajes se ajustan al nivel permitido más cercano de la rúbrica.")
    all_summary = []
    for r in reviews:
        row = {"Archivo": r["archivo"], "Estudiante": r.get("estudiante", "")}
        for s in r["secciones"]:
            row[s["Sección"]] = s["Puntaje"]
        row["Total / 10"] = r["puntaje_total"]
        all_summary.append(row)
    edited = st.data_editor(
        pd.DataFrame(all_summary), hide_index=True, use_container_width=True,
        disabled=["Archivo"] + [r["seccion"] for r in RUBRICA] + ["Total / 10"],
        key="class_editor"
    )
    # Keep source grades consistent with permitted rubric levels; this UI's grade fields are read-only
    # until section-level adjustment is implemented below.
    for r in reviews:
        with st.expander(f'{r["archivo"]} — propuesta {r["puntaje_total"]:.2f}/10', expanded=False):
            st.write(r["observaciones_generales"])
            confirmed_total = 0
            for j, s in enumerate(r["secciones"]):
                rubric = next(x for x in RUBRICA if x["seccion"] == s["Sección"])
                valid = sorted([score for score, _ in rubric["niveles"]], reverse=True)
                c1, c2, c3 = st.columns([3, 1, 4])
                with c1:
                    st.markdown(f'**{s["Sección"]}**')
                    st.caption(f'Máximo: {s["Máximo"]:.2f}')
                with c2:
                    val = st.selectbox(
                        "Puntaje confirmado", options=valid,
                        index=valid.index(s["Puntaje"]) if s["Puntaje"] in valid else 0,
                        key=f'confirmed_{reviews.index(r)}_{j}',
                        format_func=lambda x: f'{x:.2f}'
                    )
                with c3:
                    st.write(f'**Evidencia:** {s["Evidencia identificada"]}')
                    st.write(f'**Mejora:** {s["Mejora sugerida"]}')
                s["Puntaje"] = float(val)
                confirmed_total += float(val)
            r["puntaje_total"] = round(confirmed_total, 2)
            r["estudiante"] = str(edited.loc[reviews.index(r), "Estudiante"])
            st.metric("Total confirmado", f'{r["puntaje_total"]:.2f} / 10.00')
            if r.get("alertas"):
                st.markdown("**Alertas de revisión manual**")
                for a in r["alertas"]:
                    st.write(f"- {a}")
    # Recompute totals after any section edits and generate fresh workbook.
    for r in reviews:
        r["puntaje_total"] = round(sum(s["Puntaje"] for s in r["secciones"]), 2)
    st.subheader("Descargar Excel")
    st.download_button(
        "Descargar Excel consolidado (.xlsx)", data=make_excel(reviews),
        file_name="Consolidado_revision_MIN328_2026-2.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary"
    )
    with st.expander("Ver tabla detallada de todas las secciones"):
        flat = []
        for r in reviews:
            for s in r["secciones"]:
                flat.append({"Estudiante":r.get("estudiante",""),"Archivo":r["archivo"],**s})
        st.dataframe(pd.DataFrame(flat), hide_index=True, use_container_width=True)

st.caption("Privacidad: el contenido de los informes se envía al proveedor del modelo para el análisis. No cargues datos sensibles sin autorización. La extracción de texto no siempre permite verificar imágenes, paginación ni formato; confirma esos elementos manualmente.")
